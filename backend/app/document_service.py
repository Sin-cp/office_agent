"""Deterministic document tools, reusable by a future Agent.

Only direct body paragraphs are editable. Preserve all ZIP members except
word/document.xml; replace the selected paragraph's text, never global matches.
"""
import hashlib
import io
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


def read_document(path: Path):
    data = path.read_bytes()
    document = Document(io.BytesIO(data))
    return data, document, hashlib.sha256(data).hexdigest()


def heading_level(paragraph):
    properties = paragraph._p.pPr
    outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
    if outline is not None:
        value = int(outline.get(qn('w:val')))
        return value + 1 if value < 9 else None
    style = paragraph.style
    visited = set()
    while style is not None and style.style_id not in visited:
        visited.add(style.style_id)
        properties = style.element.find(qn('w:pPr'))
        outline = properties.find(qn('w:outlineLvl')) if properties is not None else None
        if outline is not None:
            value = int(outline.get(qn('w:val')))
            return value + 1 if value < 9 else None
        style = style.base_style
    return None


def editable(paragraph):
    # Reject fields, drawings, hyperlinks, revision markers and other structures.
    allowed = {qn('w:pPr'), qn('w:r')}
    if any(child.tag not in allowed for child in paragraph._p):
        return False
    return all(child.tag in {qn('w:rPr'), qn('w:t'), qn('w:tab'), qn('w:br'), qn('w:cr')}
               and (child.tag != qn('w:br') or child.get(qn('w:type'), 'textWrapping') == 'textWrapping')
               for run in paragraph.runs for child in run._r)


def blocks(document, file_id):
    result = []
    for index, paragraph in enumerate(document.paragraphs):
        level = heading_level(paragraph)
        can_edit = editable(paragraph)
        result.append(dict(id=f'{file_id}:p:{index}', position=index + 1,
                           kind='heading' if level else 'paragraph', level=level,
                           text=paragraph.text, style=paragraph.style.name if paragraph.style else '',
                           editable=can_edit,
                           reason='' if can_edit else '包含链接、图片、域或其他复杂结构，仅可查看'))
    return result


def write_replacement(data, document, paragraph, text, destination):
    # Keep paragraph properties and run formatting. New text inherits first run.
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(text)
    xml = etree.tostring(document.element, encoding='UTF-8', xml_declaration=True, standalone=True)
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(destination, 'w') as output:
        for info in source.infolist():
            output.writestr(info, xml if info.filename == 'word/document.xml' else source.read(info.filename))
