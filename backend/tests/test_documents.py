import io
import sqlite3
import zipfile
from unittest.mock import patch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import unittest
import test_files


def example():
    doc = Document()
    doc.add_heading('项目背景', level=1)
    p = doc.add_paragraph()
    p.add_run('重复文字').bold = True
    p.add_run('内容').italic = True
    doc.add_heading('实施方案', level=2)
    doc.add_paragraph('重复文字内容')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '保留表格'
    p = doc.add_paragraph('含域段落')
    field = OxmlElement('w:fldChar')
    field.set(qn('w:fldCharType'), 'begin')
    p.add_run()._r.append(field)
    doc.sections[0].header.paragraphs[0].text = '保留页眉'
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class DocumentTests(unittest.TestCase):
    setUp = test_files.FilesTest.setUp
    tearDown = test_files.FilesTest.tearDown
    upload = test_files.FilesTest.upload

    def prepare(self):
        original = example()
        file = self.upload(original).json()
        response = self.client.get(f"/api/files/{file['id']}/content")
        self.assertEqual(response.status_code, 200, response.text)
        return original, file, response.json()

    def change(self, file, content, index=2, **overrides):
        block = content['blocks'][index]
        payload = dict(block_id=block['id'], revision=content['revision'], expected_text=block['text'], new_text='测试')
        payload.update(overrides)
        return self.client.post(f"/api/files/{file['id']}/edits", json=payload)

    def test_edit_preserves_original_and_other_parts(self):
        original, file, content = self.prepare()
        self.assertEqual([b['level'] for b in content['blocks'][:3]], [1, None, 2])
        self.assertEqual(content['table_count'], 1)
        response = self.change(file, content)
        self.assertEqual(response.status_code, 201, response.text)
        result = response.json()
        self.assertEqual(result['source_file_id'], file['id'])
        new_bytes = self.client.get(f"/api/files/{result['id']}/download").content
        self.assertEqual(self.client.get(f"/api/files/{file['id']}/download").content, original)
        doc = Document(io.BytesIO(new_bytes))
        self.assertEqual(doc.paragraphs[2].text, '测试')
        self.assertEqual(doc.paragraphs[2].style.name, 'Heading 2')
        self.assertEqual(doc.paragraphs[1].text, '重复文字内容')
        self.assertEqual(doc.tables[0].cell(0, 0).text, '保留表格')
        with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(io.BytesIO(new_bytes)) as after:
            self.assertEqual(before.namelist(), after.namelist())
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(before.read(name), after.read(name))
        next_content = self.client.get(f"/api/files/{result['id']}/content").json()
        self.assertEqual(self.change(result, next_content, new_text='继续修改').status_code, 201)

    def test_duplicate_text_only_selected_block_and_first_run_format(self):
        _, file, content = self.prepare()
        response = self.change(file, content, index=1, new_text='新文字\n换行')
        doc = Document(io.BytesIO(self.client.get(f"/api/files/{response.json()['id']}/download").content))
        self.assertEqual(doc.paragraphs[1].text, '新文字\n换行')
        self.assertTrue(doc.paragraphs[1].runs[0].bold)
        self.assertEqual(doc.paragraphs[3].text, '重复文字内容')

    def test_conflicts_and_unsupported(self):
        _, file, content = self.prepare()
        for changes, status in [({'revision': 'stale'}, 409), ({'block_id': 'another:p:2'}, 404), ({'expected_text': 'wrong'}, 409), ({'new_text': '实施方案'}, 400), ({'new_text': '\u0000'}, 422)]:
            self.assertEqual(self.change(file, content, **changes).status_code, status)
        self.assertFalse(content['blocks'][4]['editable'])
        self.assertEqual(self.change(file, content, index=4).status_code, 422)
        self.assertEqual(self.client.get('/api/files').json()['total'], 1)

    def test_generation_database_failure_cleans_file(self):
        _, file, content = self.prepare()
        from pathlib import Path
        root = Path(self.temp.name) / 'uploads'
        before = set(root.iterdir())
        with self.assertLogs('app.documents', level='ERROR'), patch('app.database.insert', side_effect=sqlite3.OperationalError('test')):
            self.assertEqual(self.change(file, content).status_code, 500)
        self.assertEqual(set(root.iterdir()), before)


class MigrationTest(unittest.TestCase):
    def test_existing_database_keeps_records(self):
        import tempfile
        from pathlib import Path
        from app import database
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with sqlite3.connect(root / 'office_agent.db') as db:
                db.execute('CREATE TABLE files (id TEXT PRIMARY KEY, name TEXT NOT NULL, size INTEGER NOT NULL, storage_key TEXT NOT NULL UNIQUE, created_at TEXT NOT NULL)')
                db.execute("INSERT INTO files VALUES ('old', 'old.docx', 1, 'old.docx', '2026')")
            db.close()
            database.initialize(root)
            database.initialize(root)
            self.assertEqual(database.get(root, 'old')['name'], 'old.docx')
            self.assertIsNone(database.get(root, 'old')['source_file_id'])
